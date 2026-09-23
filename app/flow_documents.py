"""Versioned document snapshots and private evidence. No OCR or AI extraction.
Small-installation mode stores binary objects transactionally in the database so
an online database backup includes the complete dossier. See AGENTS.md for scale-out.
"""
from datetime import datetime
import csv
from decimal import Decimal
from pathlib import PurePosixPath
import hashlib
import io
import json
import os
import re
import zipfile
from fastapi import HTTPException
from sqlalchemy import select, func
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .models import Store, User
from .flow_models import FileAsset, DocTemplate, Customer
from .db import today

DOC_TITLES={'contract':'车辆订购确认书','handover':'车辆交付确认单','repair_sheet':'维修项目确认单','business':'业务处理记录'}
DOC_TITLES['retail_quote']='精品套餐报价确认单'
DEFAULT_CLAUSES={
'contract': '本单用于确认客户所选车辆、约定车辆价款及预计交付安排。\n精品加装、保险及代办服务如有选择，按关联服务单分别确认内容与金额，不重复包含在车辆价款中。\n付款安排：________________________________________\n车辆交付条件与地点：________________________________\n变更、退订及退款约定：________________________________\n其他约定：________________________________________',
'handover':'交接时请双方核对车辆外观、随车物品、钥匙及交付资料。\n随车物品：________________________________________\n现场检查与差异：____________________________________\n其他交接事项：______________________________________',
'repair_sheet':'客户确认维修项目、约定价格及结算方。超出本单的新增项目应当另行获得客户确认后施工。\n施工项目确认：______________________________________\n取车与付款安排：____________________________________',
'business':'请核对本单记录及关联凭据。\n补充事项：________________________________________\n交接与复核意见：____________________________________'}
DEFAULT_CLAUSES['retail_quote']='本单按已选择的精品套餐版本逐项列示数量及分摊金额。实际安装、交付、收付款及退款分别以原业务经办凭据为准。\n客户确认：________________________________________\n补充约定：________________________________________'
UPLOAD_LABELS={'evidence':'业务凭据','signed_contract':'合同签回件','signed_handover':'提车签回件','receipt':'收退款凭据','invoice':'发票','authorization':'客户授权','inspection':'检测记录','procurement_contract':'采购合同与核价凭据'}
ENTITY_FIELDS=('经营主体法定名称','主体识别号','登记地址','主体资料版本编号')
MAX_BYTES=int(os.getenv('FILE_MAX_BYTES','10485760'))
CASE_LIMIT=int(os.getenv('FILE_CASE_LIMIT','100'))
STORE_QUOTA=int(os.getenv('FILE_STORE_QUOTA_MB','1024'))*1024*1024


def ensure_templates(db):
    from .tenancy import single_store
    store=single_store(db)
    for kind,title in DOC_TITLES.items():
        if not db.scalar(select(DocTemplate).where(DocTemplate.kind==kind)):
            db.add(DocTemplate(store_id=store,kind=kind,title=title,clauses=DEFAULT_CLAUSES[kind],approved=False))
    db.flush()


def can_file(user,row,asset):
    if row.kind=='member_pricing_rule':
        from .member_pricing_service import can_read as member_pricing_read
        return row.flow_version==1 and member_pricing_read(user,row)
    if row.kind=='vehicle_income':return not getattr(user,'_aggregate_scope',False) and user.role in {'admin','manager','finance','auditor'}
    if row.kind in {'insurance','addon','agency'} and row.flow_version==3 and asset.category in {'receipt','invoice','procurement_contract'}:
        return user.role in {'admin','manager','finance','auditor'}
    if (row.kind in {'retail','repair','aftercare'} or row.kind=='addon' and row.flow_version==3) and asset.category=='authorization':
        return user.role in {'admin','manager','sales','service','finance','auditor'}
    if asset.category=='procurement_contract':return user.role in {'admin','manager','finance','auditor'}
    # A generated contract contains prices and identity details. Technical/stock
    # staff use neutral business/inspection documents, not sales/financial files.
    if user.role in {'inventory','technician','reception','customer_service'}:
        return asset.category in {'evidence','inspection','authorization','business'}
    return True


def snapshot(db,row,kind,*,user=None):
    if user is None:raise HTTPException(409,'文档生成与核验需要当前获权岗位，请重新登录后办理')
    from .flow_engine import get_case
    row=get_case(db,user,row.id)
    from .business_entity_service import case_entity_snapshot
    legal=case_entity_snapshot(db,user,row)
    def with_entity(info):
        # Never serialize entity_version, application evidence, current bindings
        # or bank channels. Unknown historical snapshots remain byte-for-byte
        # identical inputs to the existing fingerprint algorithm.
        if legal['status']=='frozen':
            entity=legal['entity'];info.update(zip(ENTITY_FIELDS,(entity['legal_name'],entity['tax_identifier'],entity['registered_address'],legal['context']['revision_id'])))
        return info
    if row.kind=='order' and row.flow_version in {3,4} and kind in {'contract','handover'}:
        from .sales_quote_service import document_snapshot
        return with_entity(document_snapshot(db,row,kind))
    if kind=='retail_quote':
        from .retail_bundle_documents import snapshot as retail_snapshot
        return with_entity(retail_snapshot(db,row))
    from .models import Vehicle
    from .flow_engine import scoped_get
    customer=scoped_get(db,Customer,row.customer_id) if row.customer_id else None
    vehicle=scoped_get(db,Vehicle,row.vehicle_id) if row.vehicle_id else None
    store=db.get(Store,row.store_id)
    info={'单据编号':row.number,'门店':store.name,'客户':customer.name if customer else '',
          '联系电话':customer.phone if customer else '', '车型':vehicle.model if vehicle else row.data.get('model',''),
          '车架号':vehicle.vin if vehicle else '待配车','车牌号':row.data.get('plate',''),
          '约定金额（元）':format(Decimal(row.amount_cents)/100,'.2f'),
          '预计办理日期':row.due_date.isoformat() if row.due_date else '',
          '办理内容':row.data.get('work') or row.data.get('problem') or row.data.get('topic') or row.title}
    if kind=='contract':info['选购服务']='、'.join(label for key,label in [('addon','精品加装'),('insurance','本店保险'),('agency','代办服务')] if row.data.get(key)) or '无'
    if kind=='business':
        for key in ('约定金额（元）',):info.pop(key,None)
    return with_entity(info)


def fingerprint(info):return hashlib.sha256(json.dumps(info,ensure_ascii=False,sort_keys=True).encode()).hexdigest()


def build_docx(info,title,clauses,approved,template_version):
    doc=Document();sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.15
    for border in list(doc.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        doc.styles[name].font.name='Microsoft YaHei';doc.styles[name].font.color.rgb=RGBColor(0,0,0);doc.styles[name].font.underline=False
        doc.styles[name].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.styles['Title'].font.size=Pt(21);doc.styles['Title'].font.bold=True
    p=doc.add_paragraph(title,'Title');p.alignment=1
    if not approved:
        p=doc.add_paragraph('样式草稿：条款尚未经门店确认，不用于正式签约。');p.alignment=1
        for run in p.runs:run.font.color.rgb=RGBColor.from_string('B42318');run.bold=True
    for key in ENTITY_FIELDS:
        if key in info:doc.add_paragraph(key+'：'+str(info[key]))
    long_content=info.get('办理内容','') if len(str(info.get('办理内容','')))>200 else ''
    pairs=[(k,v) for k,v in info.items() if v!='' and k not in ENTITY_FIELDS and not (k=='办理内容' and long_content)]
    table=doc.add_table(rows=0,cols=2);table.style='Table Grid';table.autofit=False
    for column,width in zip(table.columns,(Cm(4),Cm(13))):column.width=width
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    table._tbl.tblPr.append(borders)
    for key,value in pairs:
        cells=table.add_row().cells;cells[0].width=Cm(4);cells[1].width=Cm(13)
        cells[0].text=key;cells[1].text=str(value)
        for r in cells[0].paragraphs[0].runs:r.bold=True
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    if long_content:
        p=doc.add_paragraph('办理内容');p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
        for line in str(long_content).splitlines():doc.add_paragraph(line)
    p=doc.add_paragraph('约定与确认');p.paragraph_format.space_before=Pt(12);p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
    # Templates are plain text, not Jinja or executable code.
    clause_lines=clauses.splitlines()
    for index,line in enumerate(clause_lines):
        for key,value in info.items():line=line.replace('{{'+key+'}}',str(value))
        p=doc.add_paragraph(line);p.paragraph_format.keep_together=True
        if index>=len(clause_lines)-2:p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph('客户签名：________________    门店经办人：________________')
    p.paragraph_format.space_before=Pt(14);p.paragraph_format.keep_with_next=True
    doc.add_paragraph('签署日期：________________    门店确认：__________________')
    footer=sec.footer.paragraphs[0]
    footer.text=f"{info['单据编号']}  ·  模板版本 {template_version}  ·  第 "
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    footer.add_run(' 页');footer.alignment=1
    buf=io.BytesIO();doc.save(buf);return buf.getvalue()


def quota(db,row,size):
    # A short-lived store row lock serializes quotas in PostgreSQL; SQLite uses
    # transaction snapshots + one writer. Clients receive conflict/busy on races.
    db.scalar(select(Store).where(Store.id==row.store_id).with_for_update())
    count=db.scalar(select(func.count()).select_from(FileAsset).where(FileAsset.case_id==row.id)) or 0
    total=db.scalar(select(func.coalesce(func.sum(FileAsset.size),0)).where(FileAsset.store_id==row.store_id)) or 0
    if count>=CASE_LIMIT:raise HTTPException(413,'本单文件数量已达上限，请联系管理员归档')
    if total+size>STORE_QUOTA:raise HTTPException(413,'当前门店文件空间已达上限，请联系管理员扩容')


def generate_document(db,user,row,kind):
    if kind not in DOC_TITLES:raise HTTPException(422,'不支持的单据类型')
    if kind in {'contract','handover'} and row.kind!='order':raise HTTPException(422,'此业务不能生成车辆合同或交付单')
    if kind=='repair_sheet' and row.kind!='repair':raise HTTPException(422,'请选择维修工单')
    if kind=='retail_quote' and (row.kind!='retail' or row.flow_version!=2):raise HTTPException(422,'请选择有明确套餐来源的精品订单')
    if user.role in {'inventory','technician','reception','customer_service'} and kind!='business':raise HTTPException(403,'当前岗位只能生成业务记录')
    ensure_templates(db);template=db.scalar(select(DocTemplate).where(DocTemplate.kind==kind))
    info=snapshot(db,row,kind,user=user)
    from .config import settings
    if settings.environment=='production' and template.approved and kind in {'contract','handover','repair_sheet','retail_quote'} and ENTITY_FIELDS[0] not in info:
        raise HTTPException(409,'本业务缺少创建时冻结的经营主体，不能生成正式文档；不得以集团品牌或当前设置替代')
    sig=fingerprint(info)
    existing=db.scalar(select(FileAsset).where(FileAsset.case_id==row.id,FileAsset.category==kind,FileAsset.generated.is_(True),FileAsset.source_fingerprint==sig,FileAsset.template_version==template.version).order_by(FileAsset.id.desc()))
    if existing:return existing
    if kind=='retail_quote':
        from .retail_bundle_documents import build_docx as build_retail_quote
        content=build_retail_quote(info,template.title,template.clauses,template.approved,template.version)
    elif row.kind=='order' and row.flow_version in {3,4} and kind in {'contract','handover'}:
        from .sales_quote_documents import build_docx as build_sales_quote
        content=build_sales_quote(info,template.title,template.clauses,template.approved,template.version)
    else:content=build_docx(info,template.title,template.clauses,template.approved,template.version)
    quota(db,row,len(content))
    file=FileAsset(case_id=row.id,category=kind,name=f'{row.number}-{template.title}-版本{template.version}.docx',
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',sha256=hashlib.sha256(content).hexdigest(),
        content=content,size=len(content),created_by=user.id,generated=True,template_version=template.version,template_approved=template.approved,
        source_fingerprint=sig,snapshot=info)
    from .private_files import store_content
    store_content(db,file,content)
    from .file_security import initialize_file_security
    initialize_file_security(db,user,row,file)
    return file


def verify_signed(db,row,file_id,kind,*,user=None):
    from .flow_engine import scoped_get, file_exists
    file=file_exists(db,row,file_id,'signed_contract' if kind=='contract' else 'signed_handover')
    source=scoped_get(db,FileAsset,file.source_file_id) if file.source_file_id else None
    if not source or not source.generated or source.case_id!=row.id or source.category!=kind:raise HTTPException(409,'签回件必须关联本单生成的相应文档版本')
    from .file_security import require_usable
    require_usable(db,source)
    if not source.template_approved:raise HTTPException(409,'该文档使用了尚未启用的样式草稿；请主管确认模板后重新生成并签回')
    if source.source_fingerprint!=fingerprint(snapshot(db,row,kind,user=user)):raise HTTPException(409,'业务资料已变更，此签回件不是当前版本，请重新生成并签回')
    return file


def validate_upload(name,content):
    if not content or len(content)>MAX_BYTES:raise HTTPException(413,'单个文件不能为空且不得超过10兆字节')
    if not isinstance(name,str) or len(name)>180 or re.search(r'[\\/\x00-\x1f:]',name):raise HTTPException(422,'文件名无效，请使用简短文件名')
    suffix=PurePosixPath(name).suffix.lower()
    if suffix=='.pdf':
        if not content.startswith(b'%PDF-') or b'%%EOF' not in content[-2048:]:raise HTTPException(422,'文件内容不是有效的文档格式')
        if any(x.lower() in content.lower() for x in [b'/JavaScript',b'/Launch',b'/EmbeddedFile',b'/RichMedia']):raise HTTPException(422,'不接受含脚本或内嵌文件的文档')
        return 'application/pdf'
    if suffix in {'.jpg','.jpeg','.png'}:
        from PIL import Image
        try:
            with Image.open(io.BytesIO(content)) as image:
                if image.width*image.height>40_000_000 or image.format not in ({'PNG'} if suffix=='.png' else {'JPEG'}):raise ValueError()
                image.verify()
        except Exception:raise HTTPException(422,'图片损坏、格式不符或像素过大')
        return 'image/png' if suffix=='.png' else 'image/jpeg'
    if suffix=='.csv':
        if len(content)>1_000_000:raise HTTPException(413,'CSV文件不得超过1兆字节；请按业务分批导入')
        try:
            text=content.decode('utf-8-sig')
            if '\x00' in text:raise ValueError()
            count=0
            for fields in csv.reader(io.StringIO(text,newline=''),strict=True):
                count+=1
                if count>5001 or len(fields)>200 or any(len(value)>2000 for value in fields):raise ValueError()
            if count==0:raise ValueError()
        except (UnicodeDecodeError,csv.Error,ValueError):raise HTTPException(422,'CSV格式不正确；请使用UTF-8编码、完整引号和有界的行列')
        return 'text/csv; charset=utf-8'
    if suffix=='.txt':
        try:text=content.decode('utf-8')
        except UnicodeDecodeError:raise HTTPException(422,'文本文件请保存为通用编码')
        if '\x00' in text:raise HTTPException(422,'文本含有非法内容')
        return 'text/plain; charset=utf-8'
    if suffix=='.docx':
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                infos=z.infolist();names={x.filename for x in infos}
                if len(infos)>1500 or sum(x.file_size for x in infos)>40*1024*1024 or not {'[Content_Types].xml','word/document.xml'}<=names:raise ValueError()
                for x in infos:
                    low=x.filename.lower()
                    if '..' in PurePosixPath(x.filename).parts or x.filename.startswith('/') or '\\' in x.filename or x.flag_bits&1:raise ValueError()
                    if any(w in low for w in ['vba','embeddings/','activex/']) or low.endswith('.bin'):raise ValueError()
                    if x.file_size>max(100000,200*x.compress_size):raise ValueError()
                    if low.endswith(('.xml','.rels')):
                        text=z.read(x)
                        if b'<!DOCTYPE' in text or b'<!ENTITY' in text or re.search(br'TargetMode\s*=\s*[\"\']External',text,re.I):raise ValueError()
        except Exception:raise HTTPException(422,'文档损坏或含有不允许的宏、外部链接或嵌入内容')
        return 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    raise HTTPException(422,'只接受文档、照片和文本，不接受程序、网页或压缩包')


def upload_file(db,user,row,name,content,category,source_id=None):
    from .flow_engine import scoped_get, log_event
    if category not in UPLOAD_LABELS:raise HTTPException(422,'文件用途无效')
    if user.role=='auditor':raise HTTPException(403,'审计账号不能上传业务凭据')
    if row.kind in {'repair','retail','aftercare'} and category=='authorization' and user.role not in {'admin','manager','sales','service','finance'}:raise HTTPException(403,'客户报价授权由可核对价格的业务岗位上传')
    if category in {'receipt','invoice','procurement_contract'} and user.role not in {'admin','manager','finance'}:raise HTTPException(403,'财务凭据由财务岗位上传')
    if category in {'signed_contract','signed_handover'} and user.role not in {'admin','manager','sales'}:raise HTTPException(403,'合同与提车签回件由销售岗位上传')
    if category.startswith('signed_'):
        source=scoped_get(db,FileAsset,source_id) if source_id else None
        expected='contract' if category=='signed_contract' else 'handover'
        if not source or source.case_id!=row.id or not source.generated or source.category!=expected:raise HTTPException(422,'请先选择本单生成的对应文档版本')
        from .file_security import require_usable
        require_usable(db,source)
    elif source_id is not None:raise HTTPException(422,'普通凭据不应关联签回源文件')
    media=validate_upload(name,content);digest=hashlib.sha256(content).hexdigest()
    old=db.scalar(select(FileAsset).where(FileAsset.case_id==row.id,FileAsset.sha256==digest,FileAsset.category==category,FileAsset.source_file_id==source_id))
    if old:return old
    quota(db,row,len(content))
    asset=FileAsset(case_id=row.id,category=category,name=name,media_type=media,sha256=digest,size=len(content),content=content,
        created_by=user.id,source_file_id=source_id,generated=False)
    from .private_files import store_content
    store_content(db,asset,content)
    from .file_security import initialize_file_security
    initialize_file_security(db,user,row,asset)
    log_event(db,user,row,'upload','上传'+UPLOAD_LABELS[category],row.state,{'file_id':asset.id,'name':name,'sha256':digest})
    return asset


def file_info(asset,db=None):
    from .file_security import security_info, LABELS
    result={key:getattr(asset,key) for key in ['id','case_id','category','name','size','sha256','generated','template_version','template_approved','source_file_id']}|{'created_at':asset.created_at.isoformat()+'Z','label':DOC_TITLES.get(asset.category,UPLOAD_LABELS.get(asset.category,asset.category))}
    result['security']=security_info(db,asset) if db is not None else {'state':'unscanned','label':LABELS['unscanned'],'version':0,'can_use':False}
    return result
