"""V3 quote layout: compact vehicle facts and freely paginating contractual prose."""
import io
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def build_docx(info,title,clauses,approved,template_version):
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.5);section.bottom_margin=Cm(1.5);section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
    for border in list(doc.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0);style.font.underline=False
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.1
    doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.bold=True
    p=doc.add_paragraph(title,'Title');p.alignment=1
    if not approved:
        p=doc.add_paragraph('样式草稿：条款尚未经门店确认，不用于正式签约。');p.alignment=1
        for run in p.runs:run.bold=True;run.font.color.rgb=RGBColor.from_string('B42318')
    doc.add_paragraph('本单记录本次车辆报价、实际配车与客户约定。请核对车型、车架号和报价版本后签署；配套服务另单确认。')
    for key in ('经营主体法定名称','主体识别号','登记地址','主体资料版本编号'):
        if key in info:doc.add_paragraph(key+'：'+str(info[key]))
    for line in (f"单据编号：{info['单据编号']}    报价第 {info['报价版本']} 版",
                 f"门店：{info['门店']}    客户：{info['客户']}    联系电话：{info['联系电话']}"):
        doc.add_paragraph(line)
    facts=[('品牌及车系',info['品牌 / 车系']),('订购车型',info['订购车型']+' · '+info['车型目录编码']),
           ('车型参数',info['车型参数']+' · 资料版本 '+str(info['车型资料版本'])),('车架号',info['车架号']),
           ('车辆价款（元）',info['车辆约定金额（元）']),('预计交付日期',info['预计交付日期']),('本版确认有效期',info['本版确认有效期']),('另单服务',info['另单办理服务'])]
    table=doc.add_table(rows=0,cols=2);table.autofit=False;table.style='Table Grid'
    for column,width in zip(table.columns,(Cm(3.8),Cm(13.6))):column.width=width
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    table._tbl.tblPr.append(borders)
    for label,value in facts:
        row=table.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i,(cell,text) in enumerate(zip(row.cells,(label,str(value)))):
            cell.width=(Cm(3.8),Cm(13.6))[i];cell.text=text
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for run in p.runs:
                    if i==0 or label=='车辆价款（元）':run.bold=True
    doc.add_paragraph(info['金额口径'])
    def heading(text):
        p=doc.add_paragraph(text);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(7)
        for run in p.runs:run.bold=True
    heading('本版具体约定')
    # Customer terms may be long. Keep them in body prose, not a cantSplit cell.
    for line in info['本版约定'].splitlines():doc.add_paragraph(line)
    heading('条款与确认')
    clause_lines=clauses.splitlines()
    for index,line in enumerate(clause_lines):
        for key,value in info.items():line=line.replace('{{'+key+'}}',str(value))
        p=doc.add_paragraph(line)
        if index>=len(clause_lines)-2:p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph('客户签名：________________    门店经办人：________________');p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
    doc.add_paragraph('签署日期：________________    门店确认：__________________')
    p=doc.add_paragraph('报价校验摘要 '+info['报价校验摘要'])
    for run in p.runs:run.font.size=Pt(8)
    footer=section.footer.paragraphs[0];footer.alignment=1;footer.text=f"{info['单据编号']} · 报价 {info['报价版本']} · 模板 {template_version} · 第 "
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    for run in footer.runs:run.font.size=Pt(8)
    out=io.BytesIO();doc.save(out);return out.getvalue()
