"""Each frozen component has its own printable row, never an inferred unit price."""
from decimal import Decimal
from fastapi import HTTPException
from . import retail_service as retail
from .retail_bundle_service import sale_info
from .models import Store
from .flow_models import Customer


def snapshot(db,row):
    bundle=sale_info(db,row)
    if row.kind!='retail' or row.flow_version!=2 or not bundle:
        raise HTTPException(422,'精品套餐报价确认单须从已冻结套餐的原精品订单生成')
    store=retail._one(db,Store,row.store_id);customer=retail._one(db,Customer,row.customer_id)
    money=lambda v:format(Decimal(v)/100,'.2f')
    info={'单据编号':row.number,'门店':store.name,'客户':customer.name,'联系电话':customer.phone or '',
        '报价内容':'请核对以下套餐组成、每项分摊金额及退货规则；本文件生成不代表客户已授权。',
        '套餐名称':bundle['name'],'套餐版本':f"{bundle['code']} 第 {bundle['rule_version']} 版",
        '购买套数':str(bundle['sets']),'每套成交额（元）':money(bundle['price_cents_per_set'])}
    allocations={a['line_id']:a for a in bundle['allocations']}
    for idx,line in enumerate(retail._rows(db,retail.RetailLine,case_id=row.id),1):
        allocation=allocations[line.id]
        info[f'商品 {idx}']=f"{line.sku} · {line.name}；数量 {Decimal(line.quantity_milli)/1000} {line.unit}；套餐商品分摊 {money(allocation['goods_cents'])} 元"
        if line.work_item_id:
            info[f'安装 {idx}']=f"{line.work_code} · {line.work_name}；套餐安装分摊 {money(allocation['installation_cents'])} 元"
    info['本单合计（元）']=money(row.amount_cents)
    info['退货与安装规则']=bundle['mandatory_terms']
    # Bound each table row even when a configured supplement contains long paragraphs.
    text=bundle['refund_terms']
    for idx,start in enumerate(range(0,len(text),200),1):info[f'套餐补充条款 {idx}']=text[start:start+200]
    return info


def build_docx(info,title,clauses,approved,template_version):
    """A printable quote with bounded component rows and flowing prose terms."""
    import io
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();sec=doc.sections[0]
    # The bundled default Title may inherit a decorative paragraph border.
    for border in list(doc.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0);style.font.underline=False
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.1
    doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(8)
    p=doc.add_paragraph(title,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if not approved:
        p=doc.add_paragraph('样式草稿：条款尚未经门店确认，不用于正式签约。')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:run.bold=True;run.font.color.rgb=RGBColor.from_string('B42318')
    doc.add_paragraph(info['报价内容'])
    for key in ('经营主体法定名称','主体识别号','登记地址','主体资料版本编号'):
        if key in info:doc.add_paragraph(key+'：'+str(info[key]))
    for text in (f"单据编号：{info['单据编号']}    门店：{info['门店']}",
                 f"客户：{info['客户']}    联系电话：{info['联系电话']}",
                 f"套餐：{info['套餐名称']}    {info['套餐版本']}",
                 f"购买套数：{info['购买套数']}    每套成交额：{info['每套成交额（元）']} 元"):
        doc.add_paragraph(text)
    table=doc.add_table(rows=1,cols=2);table.autofit=False
    widths=(Cm(2.2),Cm(15.2))
    for column,width in zip(table.columns,widths):column.width=width
    props=table._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in (('top',70),('bottom',70),('left',100),('right',100)):
        el=OxmlElement('w:'+side);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
    props.append(margins)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    records=[('项目','冻结数量与套餐分摊')]+[(k,v) for k,v in info.items() if k.startswith(('商品 ','安装 '))]
    for idx,record in enumerate(records):
        row=table.rows[0] if idx==0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for col,(cell,text) in enumerate(zip(row.cells,record)):
            cell.width=widths[col];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=text
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                if col==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                if idx==0:
                    p.paragraph_format.keep_with_next=True
                    for run in p.runs:run.bold=True
            if idx==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDEBF7');cell._tc.get_or_add_tcPr().append(shade)
    p=doc.add_paragraph('本单成交合计：'+info['本单合计（元）']+' 元');p.paragraph_format.space_before=Pt(6)
    for run in p.runs:run.bold=True
    def heading(text):
        p=doc.add_paragraph(text);p.paragraph_format.space_before=Pt(6);p.paragraph_format.keep_with_next=True
        for run in p.runs:run.bold=True
    heading('退货与安装费用')
    doc.add_paragraph(info['退货与安装规则'])
    for k,v in info.items():
        if k.startswith('套餐补充条款 '):doc.add_paragraph(v)
    heading('约定与确认')
    clause_lines=clauses.splitlines()
    for index,line in enumerate(clause_lines):
        for k,v in info.items():line=line.replace('{{'+k+'}}',str(v))
        p=doc.add_paragraph(line);p.paragraph_format.keep_together=True
        if index>=len(clause_lines)-2:p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph('客户签名：________________    门店经办人：________________')
    p.paragraph_format.space_before=Pt(8);p.paragraph_format.keep_with_next=True
    doc.add_paragraph('签署日期：________________    门店确认：__________________')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.text=f"{info['单据编号']} · 模板版本 {template_version} · 第 "
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    for run in footer.runs:run.font.size=Pt(9)
    buf=io.BytesIO();doc.save(buf);return buf.getvalue()
